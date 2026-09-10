from flask import Flask, render_template, request, send_file
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import load_workbook, Workbook
from werkzeug.utils import secure_filename

import os
import re
import time
import uuid


app = Flask(__name__)

UPLOAD_FOLDER = "uploads"
RESULT_FOLDER = "resultados"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(RESULT_FOLDER, exist_ok=True)

URL_ONPE = "https://consultaelectoral.onpe.gob.pe/"


# ============================================================
# VALIDAR DNI
# ============================================================

def validar_dni(dni):

    dni = str(dni).strip()

    if dni.endswith(".0"):
        dni = dni[:-2]

    valido = dni.isdigit() and len(dni) == 8

    return dni, valido


# ============================================================
# BUSCAR UN VALOR EN EL TEXTO DEVUELTO POR ONPE
# ============================================================

def buscar_valor(texto, etiquetas):

    lineas = [
        linea.strip()
        for linea in texto.splitlines()
        if linea.strip()
    ]

    for i, linea in enumerate(lineas):

        linea_lower = linea.lower()

        for etiqueta in etiquetas:

            etiqueta_lower = etiqueta.lower()

            # Ejemplo:
            # Distrito: Cusco
            if etiqueta_lower in linea_lower:

                partes = re.split(
                    r":",
                    linea,
                    maxsplit=1
                )

                if len(partes) == 2:

                    valor = partes[1].strip()

                    if valor:
                        return valor

                # Ejemplo:
                # Distrito
                # Cusco

                if i + 1 < len(lineas):
                    return lineas[i + 1]

    return "No identificado"


# ============================================================
# EXTRAER INFORMACIÓN ELECTORAL
# ============================================================

def extraer_datos(texto, dni):

    texto_lower = texto.lower()

    resultado = {
        "dni": dni,
        "valido": True,
        "miembro_mesa": "No identificado",
        "nombres": "No identificado",
        "region": "No identificado",
        "provincia": "No identificado",
        "distrito": "No identificado",
        "direccion": "No identificado"
    }

    # --------------------------------------------------------
    # MIEMBRO DE MESA
    # --------------------------------------------------------

    if (
        "no eres miembro de mesa" in texto_lower
        or "no ha sido designado" in texto_lower
        or "no fue designado" in texto_lower
    ):
        resultado["miembro_mesa"] = "NO"

    elif (
        "eres miembro de mesa" in texto_lower
        or "miembro de mesa titular" in texto_lower
        or "miembro de mesa suplente" in texto_lower
        or "ha sido designado" in texto_lower
    ):
        resultado["miembro_mesa"] = "SÍ"

    else:

        valor = buscar_valor(
            texto,
            [
                "Miembro de mesa",
                "Condición de miembro de mesa"
            ]
        )

        if valor != "No identificado":
            resultado["miembro_mesa"] = valor


    # --------------------------------------------------------
    # NOMBRES
    # --------------------------------------------------------

    resultado["nombres"] = buscar_valor(
        texto,
        [
            "Nombres y apellidos",
            "Nombre completo",
            "Nombres"
        ]
    )


    # --------------------------------------------------------
    # REGIÓN / DEPARTAMENTO
    # --------------------------------------------------------

    resultado["region"] = buscar_valor(
        texto,
        [
            "Región",
            "Departamento"
        ]
    )


    # --------------------------------------------------------
    # PROVINCIA
    # --------------------------------------------------------

    resultado["provincia"] = buscar_valor(
        texto,
        [
            "Provincia"
        ]
    )


    # --------------------------------------------------------
    # DISTRITO
    # --------------------------------------------------------

    resultado["distrito"] = buscar_valor(
        texto,
        [
            "Distrito"
        ]
    )


    # --------------------------------------------------------
    # DIRECCIÓN
    # --------------------------------------------------------

    resultado["direccion"] = buscar_valor(
        texto,
        [
            "Dirección del local",
            "Dirección",
            "Direccion"
        ]
    )

    return resultado


# ============================================================
# SCRAPING REAL
# ============================================================

def consultar_onpe(dni):

    dni, valido = validar_dni(dni)

    if not valido:

        return {
            "dni": dni,
            "valido": False,
            "miembro_mesa": "DNI inválido",
            "nombres": "",
            "region": "",
            "provincia": "",
            "distrito": "",
            "direccion": ""
        }

    driver = None

    try:

        opciones = Options()

        # Chrome visible.
        # Si ONPE presenta una verificación, el usuario
        # puede completarla manualmente.
        opciones.add_argument("--start-maximized")

        driver = webdriver.Chrome(
            options=opciones
        )

        wait = WebDriverWait(
            driver,
            60
        )

        print("\n")
        print("========================================")
        print("CONSULTANDO DNI:", dni)
        print("========================================")


        # ====================================================
        # ABRIR ONPE
        # ====================================================

        driver.get(URL_ONPE)


        # ====================================================
        # BUSCAR CAMPO DNI
        # ====================================================

        inputs = wait.until(
            EC.presence_of_all_elements_located(
                (By.TAG_NAME, "input")
            )
        )

        input_dni = None

        for campo in inputs:

            tipo = (
                campo.get_attribute("type")
                or ""
            ).lower()

            placeholder = (
                campo.get_attribute("placeholder")
                or ""
            ).lower()

            nombre = (
                campo.get_attribute("name")
                or ""
            ).lower()

            identificador = (
                campo.get_attribute("id")
                or ""
            ).lower()

            if (
                "dni" in placeholder
                or "dni" in nombre
                or "dni" in identificador
                or tipo in ["text", "number", "tel"]
            ):

                input_dni = campo
                break


        if input_dni is None:

            raise Exception(
                "No se encontró el campo para ingresar DNI."
            )


        input_dni.clear()

        input_dni.send_keys(dni)

        print("DNI ingresado correctamente.")


        # ====================================================
        # BUSCAR BOTÓN
        # ====================================================

        botones = driver.find_elements(
            By.TAG_NAME,
            "button"
        )

        boton_consulta = None

        for boton in botones:

            texto_boton = (
                boton.text
                or ""
            ).strip().lower()

            if (
                "consult" in texto_boton
                or "buscar" in texto_boton
                or "continuar" in texto_boton
                or "verificar" in texto_boton
            ):

                boton_consulta = boton
                break


        # Algunos sitios usan input submit
        if boton_consulta is None:

            submits = driver.find_elements(
                By.CSS_SELECTOR,
                "input[type='submit']"
            )

            if submits:
                boton_consulta = submits[0]


        if boton_consulta is None:

            raise Exception(
                "No se encontró el botón de consulta."
            )


        # ====================================================
        # REALIZAR CONSULTA
        # ====================================================

        driver.execute_script(
            "arguments[0].scrollIntoView({block:'center'});",
            boton_consulta
        )

        time.sleep(1)

        driver.execute_script(
            "arguments[0].click();",
            boton_consulta
        )

        print("Consulta enviada a ONPE.")


        # ====================================================
        # ESPERAR RESULTADO
        # ====================================================

        texto_antes = driver.find_element(
            By.TAG_NAME,
            "body"
        ).text

        try:

            WebDriverWait(
                driver,
                30
            ).until(

                lambda navegador:

                navegador.find_element(
                    By.TAG_NAME,
                    "body"
                ).text != texto_antes

            )

        except Exception:

            # Puede ocurrir si el sitio actualiza
            # contenido sin cambiar todo el body.
            time.sleep(5)


        body = driver.find_element(
            By.TAG_NAME,
            "body"
        )

        texto_resultado = body.text


        # ====================================================
        # GUARDAR HTML REAL
        # ====================================================

        nombre_html = (
            f"respuesta_onpe_{dni}.html"
        )

        with open(
            nombre_html,
            "w",
            encoding="utf-8"
        ) as archivo:

            archivo.write(
                driver.page_source
            )


        print("\n")
        print("========== RESPUESTA REAL ONPE ==========")
        print(texto_resultado)
        print("==========================================")
        print()
        print(
            "HTML guardado en:",
            nombre_html
        )


        # ====================================================
        # EXTRAER DATOS
        # ====================================================

        resultado = extraer_datos(
            texto_resultado,
            dni
        )

        print("\nDATOS EXTRAÍDOS")
        print("---------------------------")
        print(
            "Miembro de mesa:",
            resultado["miembro_mesa"]
        )
        print(
            "Nombres:",
            resultado["nombres"]
        )
        print(
            "Región:",
            resultado["region"]
        )
        print(
            "Provincia:",
            resultado["provincia"]
        )
        print(
            "Distrito:",
            resultado["distrito"]
        )
        print(
            "Dirección:",
            resultado["direccion"]
        )
        print("---------------------------")

        return resultado


    except Exception as error:

        print(
            "\nERROR EN SCRAPING:",
            error
        )

        return {
            "dni": dni,
            "valido": True,
            "miembro_mesa": "Error de consulta",
            "nombres": "",
            "region": "",
            "provincia": "",
            "distrito": "",
            "direccion": ""
        }


    finally:

        if driver:

            time.sleep(2)

            driver.quit()


# ============================================================
# PROCESAR VARIOS DNI
# ============================================================

def procesar_dnis(dnis):

    resultados = []

    for dni in dnis:

        dni = str(dni).strip()

        if dni:

            resultado = consultar_onpe(
                dni
            )

            resultados.append(
                resultado
            )

    return resultados


# ============================================================
# PÁGINA PRINCIPAL
# ============================================================

@app.route("/")
def index():

    return render_template(
        "index.html"
    )


# ============================================================
# CONSULTA MANUAL
# ============================================================

@app.route(
    "/consultar",
    methods=["GET", "POST"]
)
def consultar():

    if request.method == "GET":

        return render_template(
            "index.html"
        )


    texto = request.form.get(
        "dnis",
        ""
    )


    dnis = [

        dni.strip()

        for dni in texto.splitlines()

        if dni.strip()

    ]


    if not dnis:

        return render_template(

            "index.html",

            error="Debe ingresar al menos un DNI."

        )


    resultados = procesar_dnis(
        dnis
    )


    return render_template(

        "index.html",

        resultados=resultados

    )


# ============================================================
# SUBIR EXCEL
# ============================================================

@app.route(
    "/subir-excel",
    methods=["POST"]
)
def subir_excel():

    if "archivo" not in request.files:

        return render_template(

            "index.html",

            error="Seleccione un archivo Excel."

        )


    archivo = request.files["archivo"]


    if archivo.filename == "":

        return render_template(

            "index.html",

            error="Seleccione un archivo Excel."

        )


    if not archivo.filename.lower().endswith(
        ".xlsx"
    ):

        return render_template(

            "index.html",

            error="El archivo debe ser .xlsx"

        )


    try:

        nombre = secure_filename(
            archivo.filename
        )


        ruta = os.path.join(
            UPLOAD_FOLDER,
            nombre
        )


        archivo.save(
            ruta
        )


        workbook = load_workbook(

            ruta,

            read_only=True,

            data_only=True

        )


        hoja = workbook.active


        encabezados = [

            str(celda.value).strip().lower()
            if celda.value is not None
            else ""

            for celda in hoja[1]

        ]


        if "dni" not in encabezados:

            workbook.close()

            return render_template(

                "index.html",

                error=(
                    "El Excel debe contener "
                    "una columna llamada 'dni'."
                )

            )


        columna_dni = (
            encabezados.index("dni")
            + 1
        )


        dnis = []


        for fila in range(
            2,
            hoja.max_row + 1
        ):

            valor = hoja.cell(

                row=fila,

                column=columna_dni

            ).value


            if valor is not None:

                dni = str(
                    valor
                ).strip()


                if dni.endswith(".0"):

                    dni = dni[:-2]


                dnis.append(
                    dni
                )


        workbook.close()


        if not dnis:

            return render_template(

                "index.html",

                error=(
                    "El Excel no contiene DNI."
                )

            )


        resultados = procesar_dnis(
            dnis
        )


        return render_template(

            "index.html",

            resultados=resultados,

            mensaje=(
                f"Se procesaron "
                f"{len(dnis)} DNI."
            )

        )


    except Exception as error:

        return render_template(

            "index.html",

            error=(
                f"Error al procesar Excel: "
                f"{error}"
            )

        )


# ============================================================
# DESCARGAR RESULTADOS
# ============================================================

@app.route(
    "/descargar",
    methods=["POST"]
)
def descargar():

    dnis = request.form.getlist(
        "dni"
    )

    miembros = request.form.getlist(
        "miembro_mesa"
    )

    nombres = request.form.getlist(
        "nombres"
    )

    regiones = request.form.getlist(
        "region"
    )

    provincias = request.form.getlist(
        "provincia"
    )

    distritos = request.form.getlist(
        "distrito"
    )

    direcciones = request.form.getlist(
        "direccion"
    )


    if not dnis:

        return (
            "No existen resultados.",
            400
        )


    wb = Workbook()

    ws = wb.active

    ws.title = "Consulta Electoral"


    # ========================================================
    # ENCABEZADOS
    # ========================================================

    ws.append([
        "DNI",
        "Miembro de mesa",
        "Nombres",
        "Región",
        "Provincia",
        "Distrito",
        "Dirección"
    ])


    # ========================================================
    # DATOS
    # ========================================================

    for i in range(
        len(dnis)
    ):

        ws.append([
            dnis[i],
            miembros[i],
            nombres[i],
            regiones[i],
            provincias[i],
            distritos[i],
            direcciones[i]
        ])


    # ========================================================
    # ANCHO DE COLUMNAS
    # ========================================================

    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 22
    ws.column_dimensions["G"].width = 55


    nombre_archivo = (

        "consulta_electoral_"

        + uuid.uuid4().hex[:8]

        + ".xlsx"

    )


    ruta = os.path.abspath(

        os.path.join(

            RESULT_FOLDER,

            nombre_archivo

        )

    )


    wb.save(
        ruta
    )


    return send_file(

        ruta,

        as_attachment=True,

        download_name="consulta_electoral.xlsx"

    )


# ============================================================
# EJECUTAR FLASK
# ============================================================

if __name__ == "__main__":

    app.run(

        host="0.0.0.0",

        port=5000,

        debug=True,

        # Evita que Flask levante dos procesos
        # mientras Selenium trabaja.
        use_reloader=False

    )